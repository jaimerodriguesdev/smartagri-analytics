"""Gera as bases integradas e os dados do dashboard SmartAgri Analytics.

Saidas:
  - Producao_Agro.csv / .xlsx
  - Qualidade_Agro.csv / .xlsx
  - Manutencao_Agro.csv / .xlsx
  - Base_Integrada_SmartAgri.xlsx
  - dados_dashboard.json / dados_dashboard.js

As tres fontes compartilham as mesmas chaves (Data, Semana, Mes_Ano, Turno,
Setor e ID_Equipamento), permitindo filtros cruzados no dashboard.
"""

from __future__ import annotations

import json
import random
from datetime import date, datetime, timedelta
from pathlib import Path

import pandas as pd
from openpyxl.styles import Font

BASE = Path(__file__).resolve().parent
SEED = 42
ANO = 2026
INICIO = date(2026, 1, 5)
SEMANAS = 40

MESES = {
    1: "Jan", 2: "Fev", 3: "Mar", 4: "Abr", 5: "Mai", 6: "Jun",
    7: "Jul", 8: "Ago", 9: "Set", 10: "Out", 11: "Nov", 12: "Dez",
}

TURNOS = ["Manhã", "Tarde", "Noite"]
FATOR_TURNO = {"Manhã": 1.04, "Tarde": 1.00, "Noite": 0.94}

SETORES = {
    "Colheita": {
        "equipamentos": ["Colheitadeira A", "Colheitadeira B"],
        "capacidade_t_h": 9.5,
    },
    "Preparo de Solo": {
        "equipamentos": ["Trator 01", "Grade Niveladora E"],
        "capacidade_t_h": 6.0,
    },
    "Semeadura": {
        "equipamentos": ["Semeadora C", "Trator 02"],
        "capacidade_t_h": 5.0,
    },
    "Pulverização": {
        "equipamentos": ["Pulverizador B", "Trator 01"],
        "capacidade_t_h": 6.5,
    },
    "Distribuição de Insumos": {
        "equipamentos": ["Distribuidor de Insumos D", "Trator 02"],
        "capacidade_t_h": 5.5,
    },
}

CUSTO_HORA = {
    "Trator 01": 320.0,
    "Trator 02": 320.0,
    "Colheitadeira A": 450.0,
    "Colheitadeira B": 450.0,
    "Pulverizador B": 380.0,
    "Semeadora C": 300.0,
    "Distribuidor de Insumos D": 340.0,
    "Grade Niveladora E": 260.0,
}

TIPOS_FALHA = ["Mecânica", "Hidráulica", "Elétrica", "Desgaste", "Sensores"]
PESOS_FALHA = [0.34, 0.22, 0.18, 0.17, 0.09]

TIPOS_DEFEITO = [
    "Grãos Quebrados",
    "Impurezas",
    "Umidade Fora do Padrão",
    "Danos Mecânicos",
    "Contaminação",
    "Classificação Incorreta",
]
PESOS_DEFEITO = [0.28, 0.22, 0.18, 0.15, 0.09, 0.08]

# Semanas com maior incidência de paradas (usadas para gerar a correlacao).
SEMANAS_CRITICAS = {6, 12, 19, 27, 34}

# Cultura predominante por mes (calendario simplificado da safra).
CULTURA_MES = {
    1: "Soja", 2: "Soja", 3: "Milho", 4: "Milho", 5: "Milho", 6: "Milho",
    7: "Algodão", 8: "Algodão", 9: "Algodão", 10: "Soja",
}


def sazonalidade(setor: str, cultura: str, mes: int) -> float:
    if setor == "Colheita":
        if cultura == "Soja" and mes in (2, 3, 10):
            return 1.25
        if cultura == "Milho" and mes in (6, 7):
            return 1.20
        if cultura == "Algodão" and mes in (9, 10):
            return 1.18
        return 0.55
    if setor == "Preparo de Solo":
        return 1.12 if mes in (1, 3, 6, 8) else 0.88
    if setor == "Semeadura":
        return 1.14 if mes in (2, 4, 7, 9) else 0.86
    if setor == "Pulverização":
        return 1.10 if mes in (3, 4, 5, 8, 9) else 0.90
    return 1.02 if mes in (1, 2, 3, 4, 5, 6, 7, 8, 9, 10) else 0.85


def gerar(rng: random.Random):
    producao = []
    manutencao = []
    qualidade = []

    for i in range(SEMANAS):
        data_semana = INICIO + timedelta(days=7 * i)
        mes = data_semana.month
        semana = f"S{i + 1:02d}"
        mes_ano = f"{ANO}-{mes:02d}"
        rotulo_mes = f"{MESES[mes]}/{ANO}"
        cultura = CULTURA_MES.get(mes, "Soja")
        critica = (i + 1) in SEMANAS_CRITICAS

        for turno in TURNOS:
            for setor, config in SETORES.items():
                equipamento = rng.choice(config["equipamentos"])
                saz = sazonalidade(setor, cultura, mes)
                horas_disponiveis = 8.0
                utilizacao = rng.uniform(0.62, 0.90) * FATOR_TURNO[turno]
                if critica:
                    utilizacao -= rng.uniform(0.07, 0.16)
                utilizacao = max(0.42, min(0.96, utilizacao))

                eficiencia = rng.uniform(0.84, 0.97) * FATOR_TURNO[turno]
                if critica:
                    eficiencia -= rng.uniform(0.02, 0.06)
                eficiencia = max(0.64, min(0.99, eficiencia))

                volume = (
                    horas_disponiveis
                    * config["capacidade_t_h"]
                    * saz
                    * utilizacao
                    * eficiencia
                )
                volume = round(volume, 1)
                horas_trabalhadas = round(horas_disponiveis * utilizacao, 2)
                eficiencia_pct = round(eficiencia * 100, 1)

                producao.append(
                    {
                        "Data": data_semana.isoformat(),
                        "Ano": ANO,
                        "Semana": semana,
                        "Mes": mes,
                        "Mes_Ano": mes_ano,
                        "Turno": turno,
                        "Setor": setor,
                        "ID_Equipamento": equipamento,
                        "Cultura": cultura,
                        "Volume_Produzido_Ton": volume,
                        "Horas_Disponiveis": horas_disponiveis,
                        "Horas_Trabalhadas": horas_trabalhadas,
                        "Eficiencia_Pct": eficiencia_pct,
                    }
                )

                # Manutencao: eventos associados ao mesmo equipamento/semana/turno.
                horas_usadas = 0.0
                prob_parada = 0.10 + (0.30 if critica else 0.0)
                if turno == "Noite":
                    prob_parada += 0.04
                if rng.random() < prob_parada:
                    qtd_eventos = 2 if (critica and rng.random() < 0.28) else 1
                    horas_usadas = 0.0
                    for _ in range(qtd_eventos):
                        limite = max(0.5, horas_disponiveis - horas_usadas - 0.2)
                        parada = rng.uniform(0.8, 5.2) + (1.6 if critica else 0.0)
                        parada = round(min(parada, limite), 1)
                        if parada <= 0:
                            continue
                        horas_usadas += parada
                        custo = parada * CUSTO_HORA[equipamento] + rng.uniform(120.0, 900.0)
                        if critica:
                            custo *= rng.uniform(1.04, 1.16)
                        manutencao.append(
                            {
                                "Data": data_semana.isoformat(),
                                "Ano": ANO,
                                "Semana": semana,
                                "Mes": mes,
                                "Mes_Ano": mes_ano,
                                "Turno": turno,
                                "Setor": setor,
                                "ID_Equipamento": equipamento,
                                "Tipo_Falha": rng.choices(
                                    TIPOS_FALHA, weights=PESOS_FALHA, k=1
                                )[0],
                                "Horas_Parada": parada,
                                "Custo_Manutencao_BRL": round(custo, 2),
                            }
                        )

                # Qualidade: uma leitura por lote processado.
                taxa = rng.uniform(0.010, 0.019)
                if horas_usadas > 0:
                    taxa += rng.uniform(0.003, 0.010)
                if critica:
                    taxa += rng.uniform(0.002, 0.006)
                if turno == "Noite":
                    taxa += 0.0015
                defeitos_kg = round(volume * 1000 * taxa * rng.uniform(0.85, 1.15), 1)
                if defeitos_kg <= 0:
                    defeitos_kg = 0.1

                pesos = PESOS_DEFEITO[:]
                if setor == "Colheita":
                    pesos[0] += 0.08
                    pesos[3] += 0.04
                elif setor in ("Preparo de Solo", "Semeadura"):
                    pesos[1] += 0.05
                elif setor == "Pulverização":
                    pesos[4] += 0.05

                qualidade.append(
                    {
                        "Data": data_semana.isoformat(),
                        "Ano": ANO,
                        "Semana": semana,
                        "Mes": mes,
                        "Mes_Ano": mes_ano,
                        "Turno": turno,
                        "Setor": setor,
                        "ID_Equipamento": equipamento,
                        "Cultura": cultura,
                        "Tipo_Defeito": rng.choices(
                            TIPOS_DEFEITO, weights=pesos, k=1
                        )[0],
                        "Defeitos_Kg": defeitos_kg,
                        "Volume_Inspecionado_Ton": volume,
                    }
                )

    return producao, manutencao, qualidade


def validar(producao, manutencao, qualidade) -> None:
    assert len(producao) == SEMANAS * len(TURNOS) * len(SETORES)
    assert producao and manutencao and qualidade
    assert all(r["Volume_Produzido_Ton"] > 0 for r in producao)
    assert all(0 < r["Eficiencia_Pct"] <= 100 for r in producao)
    assert all(r["Horas_Parada"] > 0 for r in manutencao)
    assert all(r["Custo_Manutencao_BRL"] > 0 for r in manutencao)
    assert all(r["Defeitos_Kg"] > 0 for r in qualidade)
    # Chaves de relacionamento presentes nas tres fontes.
    chaves = lambda rows: {
        (r["Semana"], r["Turno"], r["Setor"], r["ID_Equipamento"]) for r in rows
    }
    assert chaves(producao) == chaves(qualidade)
    assert chaves(manutencao) <= chaves(producao)


def arredondar(rows):
    return [
        {k: (round(v, 2) if isinstance(v, float) else v) for k, v in row.items()}
        for row in rows
    ]


def salvar_csv_xlsx(rows, nome: str) -> pd.DataFrame:
    df = pd.DataFrame(rows)
    df.to_csv(BASE / f"{nome}.csv", index=False, encoding="utf-8-sig")
    with pd.ExcelWriter(BASE / f"{nome}.xlsx", engine="openpyxl") as writer:
        df.to_excel(writer, index=False, sheet_name=nome, startrow=1)
        ws = writer.sheets[nome]
        ws["A1"] = f"SmartAgri Analytics - {nome}"
        ws["A1"].font = Font(bold=True, size=12)
        for idx, col in enumerate(df.columns, start=1):
            largura = max(len(str(col)), int(df[col].astype(str).str.len().max()))
            ws.column_dimensions[ws.cell(row=2, column=idx).column_letter].width = min(
                largura + 2, 36
            )
        ws.freeze_panes = "A3"
    return df


def gerar_dimensoes(prod_df: pd.DataFrame) -> dict[str, pd.DataFrame]:
    datas = pd.to_datetime(sorted(prod_df["Data"].unique()))
    calendario = pd.DataFrame(
        {
            "Data": datas.strftime("%Y-%m-%d"),
            "Ano": datas.year,
            "Trimestre": datas.quarter,
            "Mes": datas.month,
            "Mes_Ano": datas.strftime("%Y-%m"),
            "Mes_Label": [f"{MESES[m]}/{a}" for m, a in zip(datas.month, datas.year)],
            "Semana": [f"S{i + 1:02d}" for i in range(len(datas))],
            "Semana_Label": [f"Semana {i + 1:02d}" for i in range(len(datas))],
            "Inicio_Semana": datas.strftime("%d/%m/%Y"),
            "Fim_Semana": (datas + pd.Timedelta(days=6)).strftime("%d/%m/%Y"),
        }
    )

    equipamentos = pd.DataFrame(
        [
            {
                "ID_Equipamento": equipamento,
                "Custo_Hora_Parada_BRL": taxa,
                "Setores_Atendidos": " · ".join(
                    setor
                    for setor, config in SETORES.items()
                    if equipamento in config["equipamentos"]
                ),
            }
            for equipamento, taxa in CUSTO_HORA.items()
        ]
    ).sort_values("ID_Equipamento")

    setores = pd.DataFrame(
        [
            {
                "Setor": setor,
                "Capacidade_T_H": config["capacidade_t_h"],
                "Equipamentos": " · ".join(config["equipamentos"]),
            }
            for setor, config in SETORES.items()
        ]
    )

    turnos = pd.DataFrame(
        [
            {
                "Turno": turno,
                "Ordem": i + 1,
                "Fator_Produtividade": fator,
                "Faixa_Horaria": faixa,
            }
            for i, (turno, fator, faixa) in enumerate(
                zip(TURNOS, [FATOR_TURNO[t] for t in TURNOS], ["06h–14h", "14h–22h", "22h–06h"])
            )
        ]
    )

    culturas = pd.DataFrame(
        {
            "Cultura": ["Soja", "Milho", "Algodão"],
            "Ciclo_Principal": ["Verão", "Outono/Inverno", "Primavera"],
        }
    )

    return {
        "Dim_Calendario": calendario,
        "Dim_Equipamentos": equipamentos,
        "Dim_Setores": setores,
        "Dim_Turnos": turnos,
        "Dim_Culturas": culturas,
    }


def salvar_base_integrada(prod_df, qual_df, manut_df) -> None:
    dimensoes = gerar_dimensoes(prod_df)
    with pd.ExcelWriter(BASE / "Base_Integrada_SmartAgri.xlsx", engine="openpyxl") as writer:
        for nome, df in (
            ("Producao_Agro", prod_df),
            ("Qualidade_Agro", qual_df),
            ("Manutencao_Agro", manut_df),
            *dimensoes.items(),
        ):
            df.to_excel(writer, index=False, sheet_name=nome[:31])
            ws = writer.sheets[nome[:31]]
            ws.freeze_panes = "A2"
            for idx, col in enumerate(df.columns, start=1):
                largura = max(len(str(col)), int(df[col].astype(str).str.len().max()))
                ws.column_dimensions[ws.cell(row=1, column=idx).column_letter].width = min(
                    largura + 2, 40
                )


def salvar_json(producao, manutencao, qualidade, prod_df, qual_df, manut_df) -> None:
    payload = {
        "meta": {
            "projeto": "SmartAgri Analytics",
            "seed": SEED,
            "ano": ANO,
            "inicio": INICIO.isoformat(),
            "fim": (INICIO + timedelta(days=7 * (SEMANAS - 1) + 6)).isoformat(),
            "semanas": SEMANAS,
            "geradoEm": datetime.now().isoformat(timespec="seconds"),
            "registros": {
                "producao": len(producao),
                "qualidade": len(qualidade),
                "manutencao": len(manutencao),
            },
            "totais": {
                "volumeTon": round(prod_df["Volume_Produzido_Ton"].sum(), 1),
                "defeitosKg": round(qual_df["Defeitos_Kg"].sum(), 1),
                "horasParada": round(manut_df["Horas_Parada"].sum(), 1),
                "custoManutencao": round(manut_df["Custo_Manutencao_BRL"].sum(), 2),
            },
        },
        "producao": producao,
        "manutencao": manutencao,
        "qualidade": qualidade,
    }
    texto = json.dumps(payload, ensure_ascii=False, separators=(",", ":"))
    (BASE / "dados_dashboard.json").write_text(texto, encoding="utf-8")
    (BASE / "dados_dashboard.js").write_text(
        "window.SMARTAGRI_DATA = " + texto + ";\n", encoding="utf-8"
    )


def main() -> None:
    rng = random.Random(SEED)
    producao, manutencao, qualidade = gerar(rng)
    validar(producao, manutencao, qualidade)

    producao = arredondar(producao)
    manutencao = arredondar(manutencao)
    qualidade = arredondar(qualidade)

    prod_df = salvar_csv_xlsx(producao, "Producao_Agro")
    qual_df = salvar_csv_xlsx(qualidade, "Qualidade_Agro")
    manut_df = salvar_csv_xlsx(manutencao, "Manutencao_Agro")
    salvar_base_integrada(prod_df, qual_df, manut_df)
    salvar_json(producao, manutencao, qualidade, prod_df, qual_df, manut_df)

    disponibilidade = 100 * (
        1
        - manut_df["Horas_Parada"].sum()
        / prod_df["Horas_Disponiveis"].sum()
    )
    taxa_defeito = 100 * (
        qual_df["Defeitos_Kg"].sum()
        / (qual_df["Volume_Inspecionado_Ton"].sum() * 1000)
    )
    print(f"Producao_Agro ....: {len(prod_df):>4} linhas | {prod_df['Volume_Produzido_Ton'].sum():>12,.1f} t")
    print(f"Qualidade_Agro ...: {len(qual_df):>4} linhas | {qual_df['Defeitos_Kg'].sum():>12,.1f} kg | taxa {taxa_defeito:.2f}%")
    print(f"Manutencao_Agro ..: {len(manut_df):>4} linhas | {manut_df['Horas_Parada'].sum():>12,.1f} h | R$ {manut_df['Custo_Manutencao_BRL'].sum():,.2f}")
    print(f"Disponibilidade ..: {disponibilidade:.2f}%")
    print(f"Arquivos gerados em: {BASE}")


if __name__ == "__main__":
    main()
